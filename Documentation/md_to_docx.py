import os
import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
ER   = os.path.join(BASE, "..", "Files", "ErDiagrams")
SEQ  = os.path.join(BASE, "..", "Files", "sequence_diagrams")
MD_FILE = os.path.join(BASE, "CapFinLoan-LLD-v2.md")
OUT  = os.path.join(BASE, "CapFinLoan_Low_Level_Overview_v2.docx")

NAVY = RGBColor(0x1A, 0x23, 0x4E)
BLUE = RGBColor(0x2E, 0x4B, 0xAA)
DARK = RGBColor(0x1C, 0x1C, 0x2E)

def h(doc, text, level=1):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    if level == 1:
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(6)
        r.font.size = Pt(16); r.font.color.rgb = NAVY
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6")
        bot.set(qn("w:space"), "1"); bot.set(qn("w:color"), "2E4BAA")
        pBdr.append(bot); pPr.append(pBdr)
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        r.font.size = Pt(13); r.font.color.rgb = BLUE
    else:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        r.font.size = Pt(11); r.font.color.rgb = NAVY

def body(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.size = Pt(10.5); r.font.color.rgb = DARK

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run(text)
    r.font.size = Pt(10.5); r.font.color.rgb = DARK

def img(doc, path, width=Inches(6.0), caption=None):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=width)
        if caption:
            c = doc.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = c.add_run(caption)
            cr.italic = True; cr.font.size = Pt(9)
    else:
        body(doc, f"[Image missing: {path}]")

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.0)
    s.left_margin = s.right_margin = Cm(2.54)

# Title Page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(70)
r = p.add_run("CapFinLoan")
r.bold = True; r.font.size = Pt(36); r.font.color.rgb = NAVY

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Low-Level Design (LLD)")
r.font.size = Pt(18); r.font.color.rgb = BLUE
doc.add_page_break()

with open(MD_FILE, 'r', encoding='utf-8') as f:
    lines = f.readlines()

in_table = False
table_rows = []

for line in lines:
    stripped = line.strip()
    if not stripped:
        if in_table:
            # Render table
            headers = [c.strip() for c in table_rows[0].strip('|').split('|')]
            t = doc.add_table(rows=len(table_rows)-1, cols=len(headers))
            t.style = "Table Grid"
            
            # header
            for i, h_txt in enumerate(headers):
                if i < len(t.rows[0].cells):
                    c = t.rows[0].cells[i]
                    c.text = h_txt
                    for run in c.paragraphs[0].runs: run.bold = True
            
            # rows
            row_idx = 1
            for tr in table_rows[2:]:
                cols = [c.strip() for c in tr.strip('|').split('|')]
                for i, c_txt in enumerate(cols):
                    if i < len(t.rows[row_idx].cells):
                        t.rows[row_idx].cells[i].text = c_txt
                row_idx += 1
            
            doc.add_paragraph()
            in_table = False
            table_rows = []
        continue

    if stripped.startswith('|'):
        in_table = True
        table_rows.append(stripped)
        continue
    
    if in_table:
        in_table = False
        table_rows = []

    if stripped.startswith('---'):
        continue

    if stripped.startswith('# '):
        h(doc, stripped[2:], 1)
    elif stripped.startswith('## '):
        h(doc, stripped[3:], 1)
        # Inject images at specific H2s
        if "1. System Overview" in stripped:
            img(doc, os.path.join(ER, "architecture_diagram.png"), caption="System Architecture")
            img(doc, os.path.join(ER, "clean_architecture.png"), caption="Clean Architecture Layers")
            img(doc, os.path.join(ER, "database_er_diagram.png"), caption="Database Entity Relationship Diagram")
        elif "8. Frontend" in stripped:
            img(doc, os.path.join(ER, "frontend_component_tree.png"), caption="Angular Component Architecture")
    elif stripped.startswith('### '):
        h(doc, stripped[4:], 2)
        # Inject images at specific H3s
        if "9.1 Signup and OTP Verification" in stripped:
            img(doc, os.path.join(SEQ, "01_otp_signup_and_login.png"), caption="OTP Sign-up and Login")
        elif "9.3 Applicant Draft and Submission Flow" in stripped:
            img(doc, os.path.join(SEQ, "02_application_submit_flow.png"), caption="Application Submit Flow")
            img(doc, os.path.join(SEQ, "app_status_lifecycle.png"), caption="Application Status Lifecycle")
        elif "9.6 Document Verification Flow" in stripped:
            img(doc, os.path.join(SEQ, "03_document_upload_flow.png"), caption="Document Upload Flow")
        elif "9.5 Admin Review and Disbursal Guard" in stripped:
            img(doc, os.path.join(SEQ, "04_admin_review_and_decision.png"), caption="Admin Review and Decision Flow")
    elif stripped.startswith('- '):
        bullet(doc, stripped[2:])
    elif re.match(r'^\d+\.\s', stripped):
        bullet(doc, stripped)
    else:
        body(doc, stripped)

doc.save(OUT)
print(f"OK Saved: {OUT}")
