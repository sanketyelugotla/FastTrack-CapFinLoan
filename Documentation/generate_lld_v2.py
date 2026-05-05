import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
ER   = os.path.join(BASE, "..", "Files", "ErDiagrams")
SEQ  = os.path.join(BASE, "..", "Files", "sequence_diagrams")
OUT  = os.path.join(BASE, "CapFinLoan_Low_Level_Overview.docx")

NAVY = RGBColor(0x1A, 0x23, 0x4E)
BLUE = RGBColor(0x2E, 0x4B, 0xAA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1C, 0x1C, 0x2E)

def set_cell_bg(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True; r.font.size = Pt(16); r.font.color.rgb = NAVY
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6")
    bot.set(qn("w:space"), "1"); bot.set(qn("w:color"), "2E4BAA")
    pBdr.append(bot); pPr.append(pBdr)

def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True; r.font.size = Pt(12); r.font.color.rgb = BLUE

def body(doc, text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = bold; r.font.size = Pt(10.5); r.font.color.rgb = DARK

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25)
    r = p.add_run(text)
    r.font.size = Pt(10); r.font.color.rgb = DARK

def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    r = p.add_run(text)
    r.font.name = 'Consolas'; r.font.size = Pt(9.5)

def img(doc, path, width=Inches(6.0), caption=None):
    if os.path.exists(path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path, width=width)
        if caption:
            c = doc.add_paragraph()
            c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = c.add_run(caption)
            cr.italic = True; cr.font.size = Pt(9)
    else:
        body(doc, f"[Image missing: {path}]")

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2.0)
    s.left_margin = s.right_margin = Cm(2.54)

# COVER
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(70)
r = p.add_run("CapFinLoan")
r.bold = True; r.font.size = Pt(36); r.font.color.rgb = NAVY

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Low-Level Design (LLD)")
r.font.size = Pt(18); r.font.color.rgb = BLUE
doc.add_page_break()

# 1. System Components
h1(doc, "1. Detailed System Components")
body(doc, "The system is divided into bounded contexts (Microservices):")
h2(doc, "Auth Service")
bullet(doc, "Responsibility: User identity, registration, OTP generation, Google OAuth, and JWT issuance.")
bullet(doc, "Internal Logic: Hashes passwords with BCrypt. Validates OTPs stored in cache/DB. Signs JWT tokens with RBAC claims.")
h2(doc, "Application Service")
bullet(doc, "Responsibility: Orchestrates loan lifecycle, profile management, and Wallet ledger operations.")
bullet(doc, "Internal Logic: Calculates EMI. Validates wallet balances before state transitions (e.g., deducting fees upon submission).")
h2(doc, "Document Service")
bullet(doc, "Responsibility: Handles physical file uploads, document metadata, and verification states.")
bullet(doc, "Internal Logic: Stores files in local/cloud storage, links URLs to applications, manages status (Pending -> Verified).")
h2(doc, "Admin Service")
bullet(doc, "Responsibility: Provides review queues, analytics dashboards, and report generation.")
bullet(doc, "Internal Logic: Aggregates data from Application/Document services via events. Generates PDF reports.")
h2(doc, "Notification Service")
bullet(doc, "Responsibility: Dispatches emails and alerts asynchronously.")
bullet(doc, "Internal Logic: Consumes RabbitMQ events (e.g., ApplicationSubmittedEvent) to construct and send templates via MailKit/SMTP.")

# 2. Database Design
h1(doc, "2. Database Design (Schema)")
body(doc, "Isolated SQL Server databases per service. Key models:")
img(doc, os.path.join(ER, "database_er_diagram.png"), width=Inches(6.2), caption="Figure 2.1: Database ER Mapping")
h2(doc, "Users (AuthDb)")
code(doc, "- Id (GUID, PK)\n- FullName (NVARCHAR)\n- Email (NVARCHAR, Unique Index)\n- PasswordHash (NVARCHAR)\n- Role (NVARCHAR) -> 'APPLICANT' or 'ADMIN'")
h2(doc, "Loans (ApplicationDb)")
code(doc, "- ApplicationId (GUID, PK)\n- UserId (GUID, FK)\n- RequestedAmount (DECIMAL)\n- TenureMonths (INT)\n- EMI (DECIMAL)\n- Status (NVARCHAR, Index)")
h2(doc, "Wallet / Transactions (ApplicationDb)")
code(doc, "- TransactionId (GUID, PK)\n- UserId (GUID, FK)\n- Type (NVARCHAR) -> 'CREDIT' or 'DEBIT'\n- Amount (DECIMAL)\n- Description (NVARCHAR)\n- Timestamp (DATETIME, Index)")

# 3. API Design
h1(doc, "3. API Design")
body(doc, "All APIs are exposed through the Ocelot Gateway. Example endpoints:")
h2(doc, "POST /api/applications/{id}/submit")
bullet(doc, "Method: POST")
bullet(doc, "Request Body: None (Requires JWT)")
bullet(doc, "Logic: Checks if docs are uploaded. Checks wallet balance for processing fee. Deducts fee, updates state.")
bullet(doc, "Response (Success): { 'success': true, 'message': 'Application submitted' }")
bullet(doc, "Error Cases: 400 Bad Request if wallet insufficient or docs missing.")

h2(doc, "POST /api/wallet/topup")
bullet(doc, "Request Body: { 'amount': 5000, 'source': 'Razorpay/Mock' }")
bullet(doc, "Response: { 'success': true, 'newBalance': 15000 }")

# 4. Sequence Diagrams
h1(doc, "4. Detailed Sequence Diagrams")
h2(doc, "Loan Submission & Wallet Deduction Flow")
img(doc, os.path.join(SEQ, "02_application_submit_flow.png"), width=Inches(6.0))
h2(doc, "Document Verification & Admin Review Flow")
img(doc, os.path.join(SEQ, "04_admin_review_and_decision.png"), width=Inches(6.0))

# 5. Business Logic
h1(doc, "5. Core Business Logic")
h2(doc, "EMI Calculation Engine")
body(doc, "Formula used: E = P x R x (1+R)^N / [(1+R)^N-1]")
bullet(doc, "P = Principal (RequestedAmount)")
bullet(doc, "R = Monthly Interest Rate (e.g. 10% annual = 0.10 / 12)")
bullet(doc, "N = Tenure in Months")
h2(doc, "Wallet Ledger Logic")
body(doc, "The wallet operates on a ledger system. Balance is NOT stored as a single field. It is dynamically calculated as: SUM(Credits) - SUM(Debits).")
bullet(doc, "When a fee is deducted: Insert DEBIT transaction.")
bullet(doc, "When loan is disbursed: Insert CREDIT transaction to the user's wallet.")

# 6. Class Design
h1(doc, "6. Class & Architecture Design")
img(doc, os.path.join(ER, "clean_architecture.png"), width=Inches(6.0), caption="Clean Architecture Layers applied to each Microservice")

# 7. Security
h1(doc, "7. Security Implementation")
bullet(doc, "JWT Flow: API Gateway forwards tokens. Services use shared secret to validate signature without querying Auth Service.")
bullet(doc, "Password Hashing: BCrypt with salting.")
bullet(doc, "Role-Based Access: [Authorize(Roles=\"ADMIN\")] attributes on API endpoints; Route Guards in Angular.")

# 8. Error Handling
h1(doc, "8. Error Handling")
bullet(doc, "Global Middleware: Catches uncaught exceptions and returns standard JSON: { 'success': false, 'message': '...', 'errorCode': 500 }")
bullet(doc, "Payment/Wallet Failures: Returns specific HTTP 402 (Payment Required) or explicit 400 with 'INSUFFICIENT_FUNDS' code so UI can redirect to Top-Up.")

# 9. State Management
h1(doc, "9. State Management")
img(doc, os.path.join(SEQ, "app_status_lifecycle.png"), width=Inches(6.0), caption="Loan State Machine")

# 10. Data Flow
h1(doc, "10. Inter-Service Data Flow")
bullet(doc, "Asynchronous via RabbitMQ (MassTransit).")
bullet(doc, "Example: When Admin Approves Loan in AdminService -> publishes 'LoanApprovedEvent' -> ApplicationService consumes it to trigger Wallet Disbursal -> NotificationService consumes it to email User.")

# 11. External Integrations
h1(doc, "11. External Integrations")
bullet(doc, "Mock Payment Gateway: Simulates Razorpay behavior. Generates dummy order IDs, accepts mock webhook payloads for Top-Ups.")
bullet(doc, "Email Service: MailKit / SMTP to send OTPs and decision letters.")
bullet(doc, "AI Chatbot: Groq SDK for LLM-based intelligent assistance.")

# 12. Edge Cases
h1(doc, "12. Edge Cases Handled")
bullet(doc, "Double Payment: Idempotency keys used for Wallet top-ups. Same transaction ID cannot be processed twice.")
bullet(doc, "Negative Wallet Balance: Ledger strictly checks SUM(Credits) - SUM(Debits) >= RequestedDebit in a DB transaction before allowing a deduct.")
bullet(doc, "Concurrent Status Changes: Optimistic concurrency control on Application State row.")

# 13. Performance
h1(doc, "13. Performance Considerations")
bullet(doc, "Indexing: Primary keys (GUIDs configured sequentially where possible), unique indexes on User Email, indexes on Application Status for fast queue queries.")
bullet(doc, "Async Processing: Email sending and report generation offloaded to RabbitMQ to keep HTTP requests < 200ms.")

doc.save(OUT)
print(f"OK Saved: {OUT}")
